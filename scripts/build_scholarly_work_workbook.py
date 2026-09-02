#!/usr/bin/env python3
"""Build the public, reproducible XLSX teaching fixture for issue #20."""

from __future__ import annotations

import argparse
import csv
from datetime import datetime, timezone
from pathlib import Path
import re

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.packaging.core import DocumentProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from scholarly_work_package_utils import normalize_office_package


PALETTE = {
    "navy": "17324D",
    "blue": "2F75B5",
    "pale_blue": "DCEAF7",
    "pale_gray": "F3F5F7",
    "grid": "CBD3DA",
    "ink": "1F2933",
    "muted": "5D6874",
    "white": "FFFFFF",
    "amber": "FFF2CC",
    "amber_text": "7A5200",
    "green": "E2F0D9",
    "green_text": "215E21",
    "red": "FCE4D6",
    "red_text": "9C0006",
}
FIXED_TIME = datetime(2026, 9, 2, tzinfo=timezone.utc)
THIN = Side(style="thin", color=PALETTE["grid"])


def fix_core_modified_time(payload: bytes) -> bytes:
    """Undo openpyxl's save-time metadata mutation for byte-stable fixtures."""
    return re.sub(
        rb"(<dcterms:modified[^>]*>)[^<]+(</dcterms:modified>)",
        rb"\g<1>2026-09-02T00:00:00Z\g<2>",
        payload,
        count=1,
    )


def read_rows(path: Path, delimiter: str) -> list[list[str]]:
    with path.open(encoding="utf-8", newline="") as handle:
        return list(csv.reader(handle, delimiter=delimiter))


def add_title(sheet, last_column: str, title: str, note: str) -> None:
    sheet.merge_cells(f"A1:{last_column}1")
    sheet["A1"] = title
    sheet["A1"].fill = PatternFill("solid", fgColor=PALETTE["navy"])
    sheet["A1"].font = Font(name="Aptos Display", size=16, bold=True, color=PALETTE["white"])
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 28

    sheet.merge_cells(f"A2:{last_column}2")
    sheet["A2"] = note
    sheet["A2"].fill = PatternFill("solid", fgColor=PALETTE["pale_blue"])
    sheet["A2"].font = Font(name="Aptos", size=10, italic=True, color=PALETTE["navy"])
    sheet["A2"].alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[2].height = 34
    sheet.sheet_view.showGridLines = False


def write_table(sheet, start_row: int, rows: list[list[object]], name: str) -> str:
    for row_offset, values in enumerate(rows):
        for col_offset, value in enumerate(values, start=1):
            cell = sheet.cell(start_row + row_offset, col_offset, value)
            cell.alignment = Alignment(vertical="center", wrap_text=row_offset > 0)
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
            if row_offset == 0:
                cell.fill = PatternFill("solid", fgColor=PALETTE["blue"])
                cell.font = Font(name="Aptos", bold=True, color=PALETTE["white"])
    end_row = start_row + len(rows) - 1
    end_col = len(rows[0])
    from openpyxl.utils import get_column_letter

    ref = f"A{start_row}:{get_column_letter(end_col)}{end_row}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)
    sheet.row_dimensions[start_row].height = 32
    return ref


def set_widths(sheet, widths: list[float]) -> None:
    from openpyxl.utils import get_column_letter

    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def set_text_column(sheet, column: str, first_row: int, last_row: int) -> None:
    for row in range(first_row, last_row + 1):
        sheet[f"{column}{row}"].number_format = "@"


def build_workbook(repo_root: Path) -> Path:
    sample_root = repo_root / "examples" / "scholarly-work-foundations"
    output_path = sample_root / "output" / "scholarly-data-workbook.xlsx"
    raw_rows = read_rows(sample_root / "raw" / "postcards-messy.csv", ";")
    lookup_rows = read_rows(sample_root / "raw" / "place-lookup.csv", ",")
    cleaned_rows = read_rows(sample_root / "cleaned" / "postcards-clean.csv", ",")

    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties = DocumentProperties(
        creator="Digital Humanities Handbook",
        lastModifiedBy="Digital Humanities Handbook",
        title="Scholarly-data teaching workbook",
        subject="Original teaching fixture for scholarly-work foundations",
        description="Fictional teaching data; CC BY 4.0; version 1.0.",
        keywords="Excel, Power Query, PivotTable, data cleaning, validation",
        created=FIXED_TIME,
        modified=FIXED_TIME,
    )
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"

    readme = workbook.create_sheet("README")
    raw = workbook.create_sheet("Raw")
    lookup = workbook.create_sheet("Lookup")
    cleaned = workbook.create_sheet("Cleaned")
    transform = workbook.create_sheet("Transform Log")
    pivot = workbook.create_sheet("Pivot Check")
    chart_data = workbook.create_sheet("Chart Data")
    validation = workbook.create_sheet("Validation")
    problems = workbook.create_sheet("Known Problems")

    add_title(
        readme,
        "H",
        "Scholarly-data teaching workbook",
        "Version 1.0 • Original fictional teaching data • CC BY 4.0 • 2 September 2026",
    )
    readme_rows = [
        ["Field", "Guidance"],
        ["Purpose", "Practice controlled Excel import, repeatable Power Query transformations, PivotTables, validation and transparent charting."],
        ["Research question", "How do documented normalization choices change what this small fictional catalogue allows us to compare?"],
        ["Start", "Import raw/postcards-messy.csv into a copy. Preserve Raw and record every type and exclusion decision."],
        ["Raw", "Ten unchanged source rows, including one exact duplicate and one test record."],
        ["Cleaned", "Eight accepted rows with source fields, normalized fields, status columns and validation notes."],
        ["Pivot Check", "Formula-backed expected counts. Create your own native PivotTable from Cleaned and compare it here."],
        ["Chart Data", "Formula-linked helper range and a native chart with an explicit denominator, source note and text alternative."],
        ["Validation", "Independent totals and known-value checks. Formula cells are intentionally visible and recalculate when opened."],
        ["Interface boundary", "Excel and Power Query labels vary by platform, locale (regional settings for parsing and display) and version. Outputs and checks are the stable requirement."],
        ["Identifier fixture", "Raw and Cleaned identifiers are stored as text. In your own import, set record_id to Text before any automatic numeric conversion."],
    ]
    write_table(readme, 4, readme_rows, "ReadmeGuide")
    for row in range(5, 15):
        readme[f"A{row}"].font = Font(name="Aptos", bold=True, color=PALETTE["navy"])
    set_widths(readme, [24, 92])
    readme.freeze_panes = "A4"

    add_title(
        raw,
        "I",
        "Raw — immutable import layer",
        "Do not edit these ten records. Re-import from raw/postcards-messy.csv when testing delimiter, encoding, identifier and locale settings.",
    )
    write_table(raw, 4, raw_rows, "RawRecords")
    set_text_column(raw, "A", 5, 14)
    set_widths(raw, [15, 34, 22, 24, 16, 17, 13, 17, 38])
    raw.freeze_panes = "A5"

    add_title(
        lookup,
        "D",
        "Lookup — reviewed place-name mapping",
        "The lookup preserves each source label. Reviewed grouping values support this exercise; they are not a universal authority file.",
    )
    write_table(lookup, 4, lookup_rows, "PlaceLookup")
    set_widths(lookup, [28, 25, 16, 68])
    lookup.freeze_panes = "A5"

    add_title(
        cleaned,
        "O",
        "Cleaned — documented analytical layer",
        "One defensible interpretation for the stated exercise. Keep source expressions beside normalized values and leave uncertainty visible.",
    )
    write_table(cleaned, 4, cleaned_rows, "CleanedRecords")
    set_text_column(cleaned, "A", 5, 12)
    for row in range(5, 13):
        cleaned[f"L{row}"].number_format = "0"
    object_validation = DataValidation(type="list", formula1='"photograph,postcard,unknown"')
    date_validation = DataValidation(type="list", formula1='"valid,approximate,uncertain,invalid,missing"')
    cleaned.add_data_validation(object_validation)
    cleaned.add_data_validation(date_validation)
    object_validation.add("G5:G12")
    date_validation.add("K5:K12")
    set_widths(cleaned, [14, 34, 18, 18, 25, 20, 20, 16, 15, 14, 14, 12, 23, 16, 42])
    cleaned.freeze_panes = "B5"

    add_title(
        transform,
        "F",
        "Transform log — Power Query recipe and scholarly rationale",
        "Recreate these operations in Power Query. Applied Steps make the recipe repeatable; reasons and checks make it auditable.",
    )
    transform_rows = [
        ["Step", "Operation", "Fields", "Reason", "Check", "Failure mode"],
        [1, "Import UTF-8 semicolon-delimited text", "all", "Preserve characters and columns", "9 columns; 10 data rows", "comma chosen as delimiter"],
        [2, "Set record_id to Text before load", "record_id", "Keep leading zeros", "00009 remains five characters", "automatic whole-number type"],
        [3, "Trim outer whitespace", "title, place_raw, object_type", "Remove layout noise", "00101 title/place have no outer spaces", "source file overwritten"],
        [4, "Filter explicit test identifier", "record_id", "Exclude non-collection calibration row", "TEST-01 absent; exclusion logged", "unexplained row deletion"],
        [5, "Remove one exact duplicate across all relevant fields", "all", "Prevent an identical row from being counted twice", "one 00105 remains; IDs reconciled", "using sort order to choose among differing rows"],
        [6, "Map object-type variants", "object_type", "Support stated comparison", "3 photograph, 4 postcard, 1 unknown", "blank silently dropped"],
        [7, "Merge place lookup; retain source", "place_raw", "Group reviewed variants without erasure", "Laibach source remains; grouping is Ljubljana", "unmatched rows discarded"],
        [8, "Parse views using documented convention", "views", "Create numeric measure", "00101=1204; 00103=1125; blank stays blank", "locale coercion changes values"],
        [9, "Parse dates into value, precision and status", "date_raw", "Avoid invented precision", "00106 invalid and blank ISO", "invalid date coerced"],
        [10, "Load to Cleaned table", "all", "Separate analytical layer", "8 rows and 15 fields", "raw and cleaned mixed"],
    ]
    write_table(transform, 4, transform_rows, "TransformationLog")
    set_widths(transform, [9, 42, 24, 48, 48, 42])
    transform.freeze_panes = "A5"

    add_title(
        pivot,
        "D",
        "Pivot check — expected summary",
        "Create a native PivotTable from Cleaned with object_type_normalized in Rows and record_id in Values (Count). Compare it with these formula-backed checks.",
    )
    pivot_rows = [
        ["object_type_normalized", "Formula count", "Share", "Independent manual count"],
        ["photograph", "=COUNTIF('Cleaned'!$G$5:$G$12,A5)", "=B5/$B$8", 3],
        ["postcard", "=COUNTIF('Cleaned'!$G$5:$G$12,A6)", "=B6/$B$8", 4],
        ["unknown", "=COUNTIF('Cleaned'!$G$5:$G$12,A7)", "=B7/$B$8", 1],
        ["Total", "=SUM(B5:B7)", "=SUM(C5:C7)", 8],
    ]
    write_table(pivot, 4, pivot_rows, "PivotExpected")
    for row in range(5, 9):
        pivot[f"C{row}"].number_format = "0.0%"
    for cell in pivot[8]:
        if cell.column <= 4:
            cell.fill = PatternFill("solid", fgColor=PALETTE["pale_blue"])
            cell.font = Font(name="Aptos", bold=True, color=PALETTE["navy"])
    pivot.merge_cells("A10:D10")
    pivot["A10"] = "Interpretation check: this summary describes eight retained fictional records. It does not measure a real collection. Unknown remains in the denominator."
    pivot.merge_cells("A11:D11")
    pivot["A11"] = "Refresh test: add one controlled row to a copy of the raw CSV, refresh your query and PivotTable, and confirm that exactly one category count and the total increase by one."
    for row in (10, 11):
        pivot[f"A{row}"].fill = PatternFill("solid", fgColor=PALETTE["amber"])
        pivot[f"A{row}"].font = Font(name="Aptos", color=PALETTE["amber_text"])
        pivot[f"A{row}"].alignment = Alignment(wrap_text=True, vertical="center")
        pivot.row_dimensions[row].height = 36
    set_widths(pivot, [30, 20, 16, 30])
    pivot.freeze_panes = "A5"

    add_title(
        chart_data,
        "I",
        "Chart data — formula-linked and transparent",
        "The chart reads from formulas linked to Pivot Check. Keep denominator, category definition, exclusions, source and text alternative visible.",
    )
    chart_rows = [
        ["Normalized object type", "Record count"],
        ["photograph", "='Pivot Check'!B5"],
        ["postcard", "='Pivot Check'!B6"],
        ["unknown", "='Pivot Check'!B7"],
    ]
    write_table(chart_data, 4, chart_rows, "ChartValues")
    notes = [
        "Caption: Retained fictional catalogue records by normalized object type (n = 8).",
        "Source: Digital Humanities Handbook, Fictional Postcard Catalogue Teaching Dataset, version 1.0.",
        "Transformations/exclusions: source labels mapped to two categories plus unknown; one exact duplicate and TEST-01 excluded.",
        "Text alternative: Column chart showing photograph 3, postcard 4 and unknown 1 among eight retained records; postcard is the largest category.",
    ]
    for row, note in enumerate(notes, start=9):
        chart_data.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        chart_data.cell(row, 1, note)
        chart_data.cell(row, 1).fill = PatternFill("solid", fgColor=PALETTE["pale_gray"])
        chart_data.cell(row, 1).font = Font(name="Aptos", italic=True, color=PALETTE["muted"])
        chart_data.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="center")
        chart_data.row_dimensions[row].height = 42
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "Normalized object types (n = 8)"
    chart.y_axis.title = "Retained records"
    chart.x_axis.title = "Normalized object type"
    chart.y_axis.scaling.min = 0
    chart.y_axis.scaling.max = 5
    chart.height = 8.5
    chart.width = 13
    chart.legend = None
    chart.add_data(Reference(chart_data, min_col=2, min_row=4, max_row=7), titles_from_data=True)
    chart.set_categories(Reference(chart_data, min_col=1, min_row=5, max_row=7))
    chart_data.add_chart(chart, "D4")
    set_widths(chart_data, [30, 18, 4, 16, 16, 16, 16, 16, 16])
    chart_data.freeze_panes = "A5"

    add_title(
        validation,
        "D",
        "Validation — independent checks",
        "Expected values were counted independently from raw identifiers. Formula cells point only to cleaned and summary layers.",
    )
    validation_rows = [
        ["Check", "Expected", "Calculated", "Status"],
        ["Accepted rows", 8, "=COUNTA('Cleaned'!A5:A12)", '=IF(ABS(B5-C5)<0.0001,"PASS","CHECK")'],
        ["Pivot total", 8, "=SUM('Pivot Check'!B5:B7)", '=IF(ABS(B6-C6)<0.0001,"PASS","CHECK")'],
        ["Unknown categories", 1, '=COUNTIF(\'Cleaned\'!G5:G12,"unknown")', '=IF(ABS(B7-C7)<0.0001,"PASS","CHECK")'],
        ["Leading-zero ID retained", 1, '=COUNTIF(\'Cleaned\'!A5:A12,"00009")', '=IF(ABS(B8-C8)<0.0001,"PASS","CHECK")'],
        ["Invalid dates flagged", 1, '=COUNTIF(\'Cleaned\'!K5:K12,"invalid")', '=IF(ABS(B9-C9)<0.0001,"PASS","CHECK")'],
        ["Missing view values", 1, "=COUNTBLANK('Cleaned'!L5:L12)", '=IF(ABS(B10-C10)<0.0001,"PASS","CHECK")'],
        ["Category shares sum", 1, "=SUM('Pivot Check'!C5:C7)", '=IF(ABS(B11-C11)<0.0001,"PASS","CHECK")'],
    ]
    write_table(validation, 4, validation_rows, "ValidationChecks")
    validation["B11"].number_format = "0.0%"
    validation["C11"].number_format = "0.0%"
    validation.conditional_formatting.add(
        "D5:D11",
        CellIsRule(operator="equal", formula=['"PASS"'], fill=PatternFill("solid", fgColor=PALETTE["green"])),
    )
    validation.conditional_formatting.add(
        "D5:D11",
        CellIsRule(operator="equal", formula=['"CHECK"'], fill=PatternFill("solid", fgColor=PALETTE["red"])),
    )
    set_widths(validation, [34, 18, 18, 16])
    validation.freeze_panes = "A5"

    add_title(
        problems,
        "D",
        "Known problems — deliberate traps",
        "Use this sheet to test import and cleaning decisions. Preserve evidence and record uncertainty rather than merely producing a tidy table.",
    )
    problem_rows = [
        ["Problem", "Where", "Required response", "Do not do"],
        ["Leading zeros", "record_id 00009", "Import identifier as text", "Convert it to 9"],
        ["Exact duplicate", "second 00105", "Exclude one identical row and log the rule", "Delete both rows"],
        ["Differing duplicate candidates", "same key, different fields", "Use an explicit priority/group/index or retain/exclude rule; verify retained IDs", "Assume visible sort order chooses the retained row"],
        ["Test row", "TEST-01", "Filter through an explicit step", "Delete from raw file"],
        ["Locale-sensitive numbers", "1.204 and 1 125", "Apply the documented source convention and test known values", "Trust automatic detection"],
        ["Mixed date precision", "date_raw", "Separate display, normalized value, precision and status", "Invent missing days"],
        ["Impossible date", "00106", "Flag invalid and leave ISO date blank", "Silently coerce"],
        ["Missing values", "creator, views, type, date", "Keep missingness explicit", "Replace with zero or guesses"],
        ["Historical/bilingual places", "Laibach; Koper / Capodistria", "Retain source form beside reviewed grouping", "Erase the source wording"],
    ]
    write_table(problems, 4, problem_rows, "KnownProblems")
    set_widths(problems, [32, 32, 66, 48])
    problems.freeze_panes = "A5"

    for sheet in workbook.worksheets:
        sheet.sheet_properties.tabColor = PALETTE["blue"]
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_view.zoomScale = 90

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    normalize_office_package(
        output_path,
        transforms={"docProps/core.xml": fix_core_modified_time},
    )
    print(output_path.relative_to(repo_root))
    return output_path


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--repo-root", type=Path, required=True)
    args = parser.parse_args()
    build_workbook(args.repo_root.resolve())


if __name__ == "__main__":
    main()
