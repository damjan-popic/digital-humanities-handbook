#!/usr/bin/env python3
"""Build the original DOCX teaching fixtures for issue #20.

The script is an optional reproducibility aid; using Python is not a course
requirement. Generated prose and data are part of the CC BY 4.0 sample pack.
This build code follows the repository's MIT licence.
"""

from __future__ import annotations

import argparse
from pathlib import Path
import zipfile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "0B2545"
MUTED = "5D6874"
HEADER_FILL = "E8EEF5"
LIGHT_FILL = "F4F6F9"
WHITE = "FFFFFF"
INK = "1B1F23"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGIN_DXA = 120


def set_run_font(run, size=11, color=INK, bold=None, italic=None):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, value=CELL_MARGIN_DXA):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge in ("top", "start", "bottom", "end"):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), "80" if edge in {"top", "bottom"} else str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def set_table_text(table):
    for r_index, row in enumerate(table.rows):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    set_run_font(
                        run,
                        size=9.5,
                        color=WHITE if r_index == 0 else INK,
                        bold=(r_index == 0),
                    )
            if r_index == 0:
                set_cell_shading(cell, DARK_BLUE)
    set_repeat_table_header(table.rows[0])


def add_field(paragraph, instruction, display_text):
    begin_run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    begin_run._r.append(begin)

    instr_run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    instr_run._r.append(instr)

    sep_run = paragraph.add_run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    sep_run._r.append(sep)

    result = paragraph.add_run(display_text)
    set_run_font(result)

    end_run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._r.append(end)
    return result


def add_caption(document, kind, number, description, bookmark_name, bookmark_id):
    paragraph = document.add_paragraph(style="Caption")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(7)
    label = paragraph.add_run(f"{kind} ")
    set_run_font(label, size=10, color=DARK_BLUE, bold=True)

    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), bookmark_name)
    paragraph._p.append(start)
    field_result = add_field(paragraph, f"SEQ {kind}", str(number))
    field_result.font.size = Pt(10)
    field_result.font.bold = True
    field_result.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.append(end)

    detail = paragraph.add_run(f". {description}")
    set_run_font(detail, size=10, color=DARK_BLUE)
    return paragraph


def add_cross_reference(paragraph, bookmark, display_number):
    return add_field(paragraph, f"REF {bookmark} \\h", str(display_number))


def set_update_fields(document):
    settings = document.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def set_page_start(section, value=1):
    sect_pr = section._sectPr
    pg_num = sect_pr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num)
    pg_num.set(qn("w:start"), str(value))


def set_paragraph_bottom_border(paragraph, color="D7DBE2"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def set_paragraph_callout(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), LIGHT_FILL)
    p_pr.append(shading)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), BLUE)
    borders.append(left)
    p_pr.append(borders)


def configure_styles(document):
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    specs = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in specs.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True

    caption = styles["Caption"]
    caption.font.name = "Calibri"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    caption.font.size = Pt(10)
    caption.font.italic = False
    caption.font.color.rgb = RGBColor.from_string(DARK_BLUE)


def configure_section(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def set_running_furniture(section, left_text, right_text):
    section.header.is_linked_to_previous = False
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(3)
    left = paragraph.add_run(left_text.upper())
    set_run_font(left, size=8.5, color=MUTED, bold=True)
    right = paragraph.add_run(f"    {right_text}")
    set_run_font(right, size=8.5, color=MUTED)
    set_paragraph_bottom_border(paragraph)

    section.footer.is_linked_to_previous = False
    footer = section.footer
    fpara = footer.paragraphs[0]
    fpara.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    label = fpara.add_run("Digital Humanities Handbook  |  ")
    set_run_font(label, size=8.5, color=MUTED)
    add_field(fpara, "PAGE", "1")
    slash = fpara.add_run(" / ")
    set_run_font(slash, size=8.5, color=MUTED)
    add_field(fpara, "NUMPAGES", "4")
    for run in fpara.runs:
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor.from_string(MUTED)


def add_title_cover(document, title, subtitle, label, version):
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(44)
    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(16)
    set_run_font(kicker.add_run(label.upper()), size=10, color=BLUE, bold=True)

    title_p = document.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_after = Pt(8)
    set_run_font(title_p.add_run(title), size=28, color=NAVY, bold=True)

    sub = document.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(40)
    set_run_font(sub.add_run(subtitle), size=14, color=DARK_BLUE)

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(80)
    meta.paragraph_format.space_after = Pt(4)
    set_run_font(meta.add_run(version), size=11, color=NAVY, bold=True)
    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(
        note.add_run("Original fictional teaching material — not historical evidence"),
        size=9.5,
        color=MUTED,
        italic=True,
    )


def add_toc(document):
    document.add_heading("Contents", level=1)
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(12)
    add_field(
        paragraph,
        'TOC \\o "1-3" \\h \\z \\u',
        "Update this field in Word or LibreOffice Writer to generate the table of contents.",
    )
    note = document.add_paragraph()
    note.paragraph_format.space_after = Pt(10)
    set_run_font(
        note.add_run(
            "Interface note: field-update commands and labels can change between versions. "
            "The stable check is that headings populate the contents list and its links reach the right sections."
        ),
        size=9.5,
        color=MUTED,
        italic=True,
    )


def build_figure(path):
    width, height = 1500, 850
    image = Image.new("RGB", (width, height), f"#{WHITE}")
    draw = ImageDraw.Draw(image)
    font = ImageFont.load_default(size=26)
    bold = ImageFont.load_default(size=30)
    left, top, right, bottom = 170, 100, 1400, 690
    draw.line((left, top, left, bottom), fill=f"#{NAVY}", width=4)
    draw.line((left, bottom, right, bottom), fill=f"#{NAVY}", width=4)
    values = [("Photograph", 3), ("Postcard", 4), ("Unknown", 1)]
    colors = ["2E74B5", "1F4D78", "8A6D1D"]
    bar_width = 240
    positions = [300, 650, 1000]
    for (label, value), color, x in zip(values, colors, positions):
        bar_height = value * 115
        draw.rectangle((x, bottom - bar_height, x + bar_width, bottom), fill=f"#{color}")
        draw.text((x + 105, bottom - bar_height - 44), str(value), fill=f"#{INK}", font=bold, anchor="mm")
        draw.text((x + 120, bottom + 34), label, fill=f"#{INK}", font=font, anchor="mm")
    draw.text((35, 390), "Records", fill=f"#{INK}", font=font, anchor="mm")
    draw.text((width // 2, 44), "Normalized object types (n = 8)", fill=f"#{NAVY}", font=bold, anchor="mm")
    draw.text((width // 2, 810), "Source: original fictional teaching dataset, version 1.0", fill=f"#{MUTED}", font=font, anchor="mm")
    image.save(path, dpi=(180, 180))


def set_picture_alt(paragraph, description):
    for doc_pr in paragraph._p.xpath(".//wp:docPr"):
        doc_pr.set("descr", description)
        doc_pr.set("title", "Object-type count chart")


def set_core_properties(document, title):
    properties = document.core_properties
    properties.author = "Digital Humanities Handbook"
    properties.last_modified_by = "Digital Humanities Handbook"
    properties.title = title
    properties.subject = "Original teaching fixture for scholarly-work foundations"
    properties.keywords = "scholarly writing; Zotero; Word; LibreOffice; Excel; teaching sample"
    properties.comments = "CC BY 4.0; fictional teaching data; version 1.0"


def build_structured_paper(output_path, figure_path):
    document = Document()
    configure_styles(document)
    configure_section(document.sections[0])
    document.sections[0].header.is_linked_to_previous = False
    document.sections[0].footer.is_linked_to_previous = False
    set_core_properties(document, "Normalization as an Interpretive Intervention")
    set_update_fields(document)

    add_title_cover(
        document,
        "Normalization as an Interpretive Intervention",
        "A structured paper built from a fictional postcard catalogue",
        "Scholarly-work foundations teaching paper",
        "Version 1.0  |  2 September 2026",
    )

    main = document.add_section(WD_SECTION.NEW_PAGE)
    configure_section(main)
    set_page_start(main, 1)
    set_running_furniture(main, "Scholarly-work foundations", "Structured paper sample")

    add_toc(document)
    document.add_heading("Abstract", level=1)
    document.add_paragraph(
        "This fictional teaching study asks how documented normalization decisions change what a small postcard catalogue allows a reader to find and compare. Eight retained records show that grouping variant object types and place labels can improve comparison, but only when source strings, exclusions and unresolved cases remain visible. The exercise demonstrates a reproducible scholarly method; it does not offer historical evidence about real collections."
    )
    keywords = document.add_paragraph()
    set_run_font(keywords.add_run("Keywords: "), bold=True, color=NAVY)
    set_run_font(keywords.add_run("metadata, normalization, provenance, catalogue, data cleaning"))

    document.add_heading("1. Research question and claim", level=1)
    document.add_paragraph(
        "How do documented normalization decisions change what a small fictional postcard catalogue allows a reader to find and compare? The provisional claim is that normalization improves comparison across inconsistent descriptions, while remaining defensible only when the cleaned table preserves source expressions, names exclusions and exposes uncertainty."
    )
    document.add_heading("1.1 Distinguishing description and interpretation", level=2)
    document.add_paragraph(
        "The raw table contains four source expressions for two related object types. That is a description of the teaching data. Treating those expressions as two analytical categories is an interpretation encoded as a transformation. The two statements must not be collapsed."
    )
    document.add_heading("2. Materials and method", level=1)
    document.add_paragraph(
        "The immutable raw layer contains ten records. The method excludes one exact duplicate and one identifier explicitly marked as a test, then retains eight records. It trims surrounding whitespace, preserves identifiers as text, separates source and normalized place labels, records date precision and maps object-type variants to photograph, postcard or unknown. Each decision is reproduced in the transformation log and checked against an independent validation layer."
    )
    document.add_heading("2.1 Data and provenance", level=2)
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ["Layer", "Rows", "Function", "Change policy"]
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    rows = [
        ("Raw", "10", "Preserve imported source values", "Never overwrite"),
        ("Cleaned", "8", "Apply documented analytical fields", "Refresh from raw"),
        ("Output", "8", "Summarize for a stated question", "Rebuild after changes"),
        ("Validation", "—", "Check totals and known values", "Keep independent"),
    ]
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    set_table_geometry(table, [1500, 960, 3300, 3600])
    set_table_text(table)
    add_caption(
        document,
        "Table",
        1,
        "Layer separation in the teaching data package.",
        "tbl_layers",
        1,
    )
    source = document.add_paragraph()
    source.paragraph_format.space_before = Pt(4)
    source.paragraph_format.space_after = Pt(4)
    set_run_font(
        source.add_run("Source: Digital Humanities Handbook original teaching sample, version 1.0."),
        size=9,
        color=MUTED,
        italic=True,
    )

    document.add_heading("3. Results", level=1)
    paragraph = document.add_paragraph()
    paragraph.add_run("The separation of layers summarized in Table ")
    add_cross_reference(paragraph, "tbl_layers", 1)
    paragraph.add_run(
        " makes the chart denominator recoverable. Of the eight retained records, four are normalized as postcards, three as photographs and one as unknown. Missing values remain missing rather than becoming zero."
    )
    for run in paragraph.runs:
        set_run_font(run)

    fig_paragraph = document.add_paragraph()
    fig_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fig_paragraph.paragraph_format.keep_with_next = True
    fig_paragraph.add_run().add_picture(str(figure_path), width=Inches(5.8))
    set_picture_alt(
        fig_paragraph,
        "Bar chart of eight retained records: photograph 3, postcard 4, unknown 1.",
    )
    add_caption(
        document,
        "Figure",
        1,
        "Retained records by normalized object type (n = 8).",
        "fig_types",
        2,
    )
    fig_source = document.add_paragraph()
    fig_source.paragraph_format.space_before = Pt(4)
    fig_source.paragraph_format.space_after = Pt(4)
    set_run_font(
        fig_source.add_run(
            "Source: `output/chart-data.csv`; exclusions are one exact duplicate and one test record."
        ),
        size=9,
        color=MUTED,
        italic=True,
    )
    paragraph = document.add_paragraph()
    paragraph.add_run("Figure ")
    add_cross_reference(paragraph, "fig_types", 1)
    paragraph.add_run(
        " retains the unknown category. Removing it would make the chart look cleaner while concealing unresolved metadata."
    )
    for run in paragraph.runs:
        set_run_font(run)

    document.add_heading("4. Discussion", level=1)
    document.add_heading("4.1 Claim, evidence and reasoning", level=2)
    document.add_paragraph(
        "Grouping makes inconsistent strings countable, but it does not reveal a natural category. The interpretive warrant comes from the stated research question and the preserved source fields. In particular, mapping Laibach to Ljubljana supports present-day aggregation while retaining the historical expression for inspection."
    )
    document.add_heading("4.2 Counterargument", level=2)
    document.add_paragraph(
        "A reader may object that retaining both source and normalized values complicates the table. The complication is methodologically useful: separate fields permit analysis without erasing the evidence from which the grouping was made."
    )
    document.add_heading("4.3 Limitations", level=2)
    document.add_paragraph(
        "The corpus is small, fictional and designed to expose common data problems. Its counts demonstrate an auditable workflow but cannot support a claim about real catalogues, institutions or historical patterns. Place mappings also require contextual expertise beyond string matching."
    )

    document.add_heading("5. Conclusion", level=1)
    document.add_paragraph(
        "The most useful cleaned dataset is not the smoothest one. It is the one whose transformations, exclusions and uncertainties a reader can reconstruct and contest."
    )
    document.add_heading("References", level=1)
    references = [
        "Digital Humanities Handbook. Fictional Postcard Catalogue Teaching Dataset. Version 1.0. 2026.",
        "Digital Humanities Handbook. “Reading Notes for the Fictional Postcard-Catalogue Study.” Version 1.0. 2026.",
    ]
    for text in references:
        p = document.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.first_line_indent = Inches(-0.3)
        p.paragraph_format.space_after = Pt(6)
        p.add_run(text)

    document.save(output_path)


def build_bibliography(output_path):
    document = Document()
    configure_styles(document)
    configure_section(document.sections[0])
    set_core_properties(document, "Citation-Style Audit — Illustrative Generated Output")
    set_update_fields(document)
    set_running_furniture(
        document.sections[0],
        "Citation-style audit",
        "Illustrative generated output",
    )

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(18)
    title.paragraph_format.space_after = Pt(7)
    set_run_font(title.add_run("Citation-Style Audit"), size=24, color=NAVY, bold=True)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(18)
    set_run_font(
        subtitle.add_run("Illustrative generated notes and bibliography"),
        size=13,
        color=DARK_BLUE,
    )

    callout = document.add_paragraph()
    callout.paragraph_format.space_before = Pt(6)
    callout.paragraph_format.space_after = Pt(12)
    callout.add_run(
        "STATIC TEACHING OUTPUT. Regenerate this document from Zotero, record the exact local CSL label/version, "
        "and compare every field with the current governing authority. Do not hand-correct output that originates in bad metadata."
    )
    set_paragraph_callout(callout)
    for run in callout.runs:
        set_run_font(run, size=10, color=NAVY, bold=True)

    document.add_heading("Audit basis", level=1)
    document.add_paragraph(
        "System: Chicago notes and bibliography as represented by the current Chicago Manual of Style Online sample-citation guide. Authority and Zotero style repository accessed 2 September 2026. This file demonstrates the outputs to inspect; it does not freeze a mutable CSL file."
    )

    document.add_heading("Full note", level=1)
    full = document.add_paragraph()
    full.paragraph_format.first_line_indent = Inches(0.3)
    full.add_run("1. Eric Hayot, ")
    title_run = full.add_run("The Elements of Academic Style: Writing for the Humanities")
    title_run.italic = True
    full.add_run(" (New York: Columbia University Press, 2014), 41–43.")

    document.add_heading("Shortened repeat note", level=1)
    short = document.add_paragraph()
    short.paragraph_format.first_line_indent = Inches(0.3)
    short.add_run("2. Hayot, ")
    title_run = short.add_run("Elements of Academic Style")
    title_run.italic = True
    short.add_run(", 45.")

    document.add_heading("Bibliography", level=1)
    entries = [
        ("Digital Humanities Handbook.", "Fictional Postcard Catalogue Teaching Dataset", ". Version 1.0. 2026. https://github.com/damjan-popic/digital-humanities-handbook."),
        ("Digital Humanities Handbook.", "Scholarly-Data Teaching Workbook", ". Version 1.0. 2026. https://github.com/damjan-popic/digital-humanities-handbook."),
        ("Gump, Steven E.", "The Elements of Academic Style: Writing for the Humanities by Eric Hayot (review)", ". Journal of Scholarly Publishing 46, no. 4 (2015): 399–403. https://doi.org/10.3138/jsp.46.4.BR2."),
        ("Harvard College Writing Center.", "Developing A Thesis", ". Accessed September 2, 2026. https://writingcenter.fas.harvard.edu/thesis."),
        ("Hayot, Eric.", "The Elements of Academic Style: Writing for the Humanities", ". New York: Columbia University Press, 2014."),
    ]
    for creator, work, remainder in entries:
        p = document.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.first_line_indent = Inches(-0.3)
        p.paragraph_format.space_after = Pt(7)
        p.add_run(f"{creator} ")
        work_run = p.add_run(work)
        work_run.italic = True
        p.add_run(remainder)

    document.add_heading("Field-level checks", level=1)
    checks = document.add_table(rows=1, cols=3)
    checks.style = "Table Grid"
    for cell, value in zip(checks.rows[0].cells, ["Output", "Check", "If wrong"]):
        cell.text = value
    data = [
        ("Book", "creator, italic title, place, publisher, year", "correct Zotero item fields"),
        ("Review", "reviewed-work wording, journal, issue, pages, DOI", "verify type and DOI record"),
        ("Webpage", "corporate author, title, URL, access date", "verify live landing page"),
        ("Local files", "corporate creator, title, version, repository", "record editorial decision"),
    ]
    for values in data:
        cells = checks.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    set_table_geometry(checks, [1600, 4300, 3460])
    set_table_text(checks)

    document.add_heading("Unresolved case", level=1)
    document.add_paragraph(
        "The workbook may be treated as software, a dataset or a supporting file depending on its publication context. The model records the ambiguity instead of pretending that a generic guide supplies a universal answer."
    )
    document.save(output_path)


def build_odt(output_path, figure_path):
    """Create a valid structured ODT counterpart without requiring LibreOffice."""
    mimetype = "application/vnd.oasis.opendocument.text"
    content = """<?xml version="1.0" encoding="UTF-8"?>
<office:document-content
 xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
 xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0"
 xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"
 xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0"
 xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0"
 xmlns:xlink="http://www.w3.org/1999/xlink"
 xmlns:svg="urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0"
 office:version="1.3">
 <office:automatic-styles>
  <style:style style:name="CoverTitle" style:family="paragraph">
   <style:paragraph-properties fo:text-align="center" fo:margin-top="2.0in" fo:margin-bottom="0.10in"/>
   <style:text-properties style:font-name="Calibri" fo:font-size="28pt" fo:font-weight="bold" fo:color="#0B2545"/>
  </style:style>
  <style:style style:name="CoverSubtitle" style:family="paragraph">
   <style:paragraph-properties fo:text-align="center" fo:margin-bottom="0.50in"/>
   <style:text-properties style:font-name="Calibri" fo:font-size="14pt" fo:color="#1F4D78"/>
  </style:style>
  <style:style style:name="PageBreak" style:family="paragraph"><style:paragraph-properties fo:break-before="page"/></style:style>
  <style:style style:name="Caption" style:family="paragraph">
   <style:paragraph-properties fo:margin-top="0.06in" fo:margin-bottom="0.10in"/>
   <style:text-properties style:font-name="Calibri" fo:font-size="10pt" fo:color="#1F4D78"/>
  </style:style>
  <style:style style:name="Source" style:family="paragraph">
   <style:paragraph-properties fo:margin-top="0.06in" fo:margin-bottom="0.06in"/>
   <style:text-properties style:font-name="Calibri" fo:font-size="9pt" fo:font-style="italic" fo:color="#5D6874"/>
  </style:style>
  <style:style style:name="Figure" style:family="graphic">
   <style:graphic-properties style:horizontal-pos="center" style:horizontal-rel="paragraph" style:wrap="none"/>
  </style:style>
  <style:style style:name="Table1" style:family="table"><style:table-properties style:width="6.5in" table:align="margins"/></style:style>
  <style:style style:name="TableCol1" style:family="table-column"><style:table-column-properties style:column-width="1.05in"/></style:style>
  <style:style style:name="TableCol2" style:family="table-column"><style:table-column-properties style:column-width="0.65in"/></style:style>
  <style:style style:name="TableCol3" style:family="table-column"><style:table-column-properties style:column-width="2.30in"/></style:style>
  <style:style style:name="TableCol4" style:family="table-column"><style:table-column-properties style:column-width="2.50in"/></style:style>
  <style:style style:name="HeaderCell" style:family="table-cell"><style:table-cell-properties fo:background-color="#1F4D78" fo:border="0.5pt solid #FFFFFF" fo:padding="0.08in"/><style:text-properties fo:color="#FFFFFF" fo:font-weight="bold"/></style:style>
  <style:style style:name="BodyCell" style:family="table-cell"><style:table-cell-properties fo:border="0.5pt solid #AAB3BC" fo:padding="0.08in"/></style:style>
 </office:automatic-styles>
 <office:body>
  <office:text>
   <text:sequence-decls><text:sequence-decl text:display-outline-level="0" text:name="Table"/><text:sequence-decl text:display-outline-level="0" text:name="Figure"/></text:sequence-decls>
   <text:p text:style-name="CoverTitle">Normalization as an Interpretive Intervention</text:p>
   <text:p text:style-name="CoverSubtitle">A structured paper built from a fictional postcard catalogue</text:p>
   <text:p text:style-name="CoverSubtitle">Version 1.0 | 2 September 2026</text:p>
   <text:p text:style-name="Source">Original fictional teaching material — not historical evidence</text:p>
   <text:h text:outline-level="1" text:style-name="PageBreak">Contents</text:h>
   <text:table-of-content text:name="Table of Contents1">
    <text:table-of-content-source text:outline-level="3" text:use-outline-level="true"/>
    <text:index-body><text:p>Update this table of contents in LibreOffice Writer.</text:p></text:index-body>
   </text:table-of-content>
   <text:h text:outline-level="1">Abstract</text:h>
   <text:p>This fictional teaching study asks how documented normalization decisions change what a small postcard catalogue allows a reader to find and compare. Eight retained records show that grouping variant object types and place labels can improve comparison, but only when source strings, exclusions and unresolved cases remain visible. The exercise demonstrates a reproducible scholarly method; it does not offer historical evidence about real collections.</text:p>
   <text:h text:outline-level="1">1. Research question and claim</text:h>
   <text:p>How do documented normalization decisions change what a small fictional postcard catalogue allows a reader to find and compare? The provisional claim is that normalization improves comparison across inconsistent descriptions, while remaining defensible only when the cleaned table preserves source expressions, names exclusions and exposes uncertainty.</text:p>
   <text:h text:outline-level="2">1.1 Distinguishing description and interpretation</text:h>
   <text:p>The raw table contains four source expressions for two related object types. That is a description of the teaching data. Treating those expressions as two analytical categories is an interpretation encoded as a transformation.</text:p>
   <text:h text:outline-level="1">2. Materials and method</text:h>
   <text:p>The immutable raw layer contains ten records. The method excludes one exact duplicate and one identifier marked as a test, then retains eight records. It trims surrounding whitespace, preserves identifiers as text, separates source and normalized place labels, records date precision and maps object-type variants to photograph, postcard or unknown.</text:p>
   <text:h text:outline-level="2">2.1 Data and provenance</text:h>
   <table:table table:name="LayerTable" table:style-name="Table1">
    <table:table-column table:style-name="TableCol1"/><table:table-column table:style-name="TableCol2"/><table:table-column table:style-name="TableCol3"/><table:table-column table:style-name="TableCol4"/>
    <table:table-header-rows><table:table-row>
     <table:table-cell table:style-name="HeaderCell" office:value-type="string"><text:p>Layer</text:p></table:table-cell>
     <table:table-cell table:style-name="HeaderCell" office:value-type="string"><text:p>Rows</text:p></table:table-cell>
     <table:table-cell table:style-name="HeaderCell" office:value-type="string"><text:p>Function</text:p></table:table-cell>
     <table:table-cell table:style-name="HeaderCell" office:value-type="string"><text:p>Change policy</text:p></table:table-cell>
    </table:table-row></table:table-header-rows>
    <table:table-row><table:table-cell table:style-name="BodyCell" office:value-type="string"><text:p>Raw</text:p></table:table-cell><table:table-cell table:style-name="BodyCell" office:value-type="float" office:value="10"><text:p>10</text:p></table:table-cell><table:table-cell table:style-name="BodyCell" office:value-type="string"><text:p>Preserve imported source values</text:p></table:table-cell><table:table-cell table:style-name="BodyCell" office:value-type="string"><text:p>Never overwrite</text:p></table:table-cell></table:table-row>
    <table:table-row><table:table-cell table:style-name="BodyCell" office:value-type="string"><text:p>Cleaned</text:p></table:table-cell><table:table-cell table:style-name="BodyCell" office:value-type="float" office:value="8"><text:p>8</text:p></table:table-cell><table:table-cell table:style-name="BodyCell" office:value-type="string"><text:p>Apply documented analytical fields</text:p></table:table-cell><table:table-cell table:style-name="BodyCell" office:value-type="string"><text:p>Refresh from raw</text:p></table:table-cell></table:table-row>
    <table:table-row><table:table-cell table:style-name="BodyCell" office:value-type="string"><text:p>Output</text:p></table:table-cell><table:table-cell table:style-name="BodyCell" office:value-type="float" office:value="8"><text:p>8</text:p></table:table-cell><table:table-cell table:style-name="BodyCell" office:value-type="string"><text:p>Summarize for a stated question</text:p></table:table-cell><table:table-cell table:style-name="BodyCell" office:value-type="string"><text:p>Rebuild after changes</text:p></table:table-cell></table:table-row>
    <table:table-row><table:table-cell table:style-name="BodyCell" office:value-type="string"><text:p>Validation</text:p></table:table-cell><table:table-cell table:style-name="BodyCell" office:value-type="string"><text:p>—</text:p></table:table-cell><table:table-cell table:style-name="BodyCell" office:value-type="string"><text:p>Check totals and known values</text:p></table:table-cell><table:table-cell table:style-name="BodyCell" office:value-type="string"><text:p>Keep independent</text:p></table:table-cell></table:table-row>
   </table:table>
   <text:p text:style-name="Caption"><text:reference-mark-start text:name="tbl_layers"/>Table <text:sequence text:name="Table" text:formula="ooow:Table+1">1</text:sequence>. Layer separation in the teaching data package.<text:reference-mark-end text:name="tbl_layers"/></text:p>
   <text:p text:style-name="Source">Source: Digital Humanities Handbook original teaching sample, version 1.0.</text:p>
   <text:h text:outline-level="1">3. Results</text:h>
   <text:p>The separation of layers summarized in <text:reference-ref text:ref-name="tbl_layers" text:reference-format="text">Table 1</text:reference-ref> makes the chart denominator recoverable. Of the eight retained records, four are normalized as postcards, three as photographs and one as unknown.</text:p>
   <text:p><draw:frame draw:style-name="Figure" draw:name="Object-type chart" text:anchor-type="paragraph" svg:width="5.8in" svg:height="3.29in"><draw:image xlink:href="Pictures/object-type-chart.png" xlink:type="simple" xlink:show="embed" xlink:actuate="onLoad" draw:mime-type="image/png"/></draw:frame></text:p>
   <text:p text:style-name="Caption"><text:reference-mark-start text:name="fig_types"/>Figure <text:sequence text:name="Figure" text:formula="ooow:Figure+1">1</text:sequence>. Retained records by normalized object type (n = 8).<text:reference-mark-end text:name="fig_types"/></text:p>
   <text:p text:style-name="Source">Source: output/chart-data.csv; exclusions are one exact duplicate and one test record.</text:p>
   <text:p><text:reference-ref text:ref-name="fig_types" text:reference-format="text">Figure 1</text:reference-ref> retains the unknown category. Removing it would make the chart look cleaner while concealing unresolved metadata.</text:p>
   <text:h text:outline-level="1">4. Discussion</text:h>
   <text:h text:outline-level="2">4.1 Claim, evidence and reasoning</text:h>
   <text:p>Grouping makes inconsistent strings countable, but it does not reveal a natural category. The interpretive warrant comes from the stated research question and the preserved source fields. Mapping Laibach to Ljubljana supports present-day aggregation while retaining the historical expression for inspection.</text:p>
   <text:h text:outline-level="2">4.2 Counterargument</text:h>
   <text:p>A reader may object that retaining both source and normalized values complicates the table. The complication is methodologically useful: separate fields permit analysis without erasing the evidence from which the grouping was made.</text:p>
   <text:h text:outline-level="2">4.3 Limitations</text:h>
   <text:p>The corpus is small, fictional and designed to expose common data problems. Its counts demonstrate an auditable workflow but cannot support a claim about real catalogues, institutions or historical patterns.</text:p>
   <text:h text:outline-level="1">5. Conclusion</text:h>
   <text:p>The most useful cleaned dataset is not the smoothest one. It is the one whose transformations, exclusions and uncertainties a reader can reconstruct and contest.</text:p>
   <text:h text:outline-level="1">References</text:h>
   <text:p>Digital Humanities Handbook. <text:span text:style-name="Emphasis">Fictional Postcard Catalogue Teaching Dataset</text:span>. Version 1.0. 2026.</text:p>
   <text:p>Digital Humanities Handbook. “Reading Notes for the Fictional Postcard-Catalogue Study.” Version 1.0. 2026.</text:p>
  </office:text>
 </office:body>
</office:document-content>"""
    # content.xml uses fo properties, so include that namespace explicitly.
    content = content.replace(
        'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"',
        'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0"',
    )
    styles = """<?xml version="1.0" encoding="UTF-8"?>
<office:document-styles xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0" xmlns:svg="urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0" office:version="1.3">
 <office:font-face-decls><style:font-face style:name="Calibri" svg:font-family="Calibri"/></office:font-face-decls>
 <office:styles>
  <style:default-style style:family="paragraph"><style:paragraph-properties fo:margin-top="0in" fo:margin-bottom="0.083in" fo:line-height="125%"/><style:text-properties style:font-name="Calibri" fo:font-size="11pt" fo:color="#1B1F23"/></style:default-style>
  <style:style style:name="Standard" style:family="paragraph" style:class="text"/>
  <style:style style:name="Heading" style:family="paragraph" style:class="text"><style:paragraph-properties fo:keep-with-next="always"/><style:text-properties style:font-name="Calibri" fo:font-weight="bold" fo:color="#2E74B5"/></style:style>
  <style:style style:name="Heading_20_1" style:display-name="Heading 1" style:family="paragraph" style:parent-style-name="Heading" style:next-style-name="Standard" style:default-outline-level="1"><style:paragraph-properties fo:margin-top="0.25in" fo:margin-bottom="0.139in"/><style:text-properties fo:font-size="16pt"/></style:style>
  <style:style style:name="Heading_20_2" style:display-name="Heading 2" style:family="paragraph" style:parent-style-name="Heading" style:next-style-name="Standard" style:default-outline-level="2"><style:paragraph-properties fo:margin-top="0.194in" fo:margin-bottom="0.097in"/><style:text-properties fo:font-size="13pt"/></style:style>
  <style:style style:name="Emphasis" style:family="text"><style:text-properties fo:font-style="italic"/></style:style>
 </office:styles>
 <office:automatic-styles>
  <style:page-layout style:name="pm1"><style:page-layout-properties fo:page-width="8.5in" fo:page-height="11in" style:print-orientation="portrait" fo:margin-top="1in" fo:margin-bottom="1in" fo:margin-left="1in" fo:margin-right="1in"/></style:page-layout>
 </office:automatic-styles>
 <office:master-styles><style:master-page style:name="Standard" style:page-layout-name="pm1"/></office:master-styles>
</office:document-styles>"""
    meta = """<?xml version="1.0" encoding="UTF-8"?>
<office:document-meta xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" xmlns:meta="urn:oasis:names:tc:opendocument:xmlns:meta:1.0" xmlns:dc="http://purl.org/dc/elements/1.1/" office:version="1.3"><office:meta><dc:title>Normalization as an Interpretive Intervention</dc:title><dc:creator>Digital Humanities Handbook</dc:creator><dc:description>Original fictional teaching fixture; CC BY 4.0.</dc:description><meta:keyword>scholarly writing</meta:keyword><meta:keyword>LibreOffice Writer</meta:keyword><meta:generator>Digital Humanities Handbook reproducible fixture builder</meta:generator></office:meta></office:document-meta>"""
    settings = """<?xml version="1.0" encoding="UTF-8"?>
<office:document-settings xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" office:version="1.3"><office:settings/></office:document-settings>"""
    manifest = """<?xml version="1.0" encoding="UTF-8"?>
<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.3">
 <manifest:file-entry manifest:full-path="/" manifest:media-type="application/vnd.oasis.opendocument.text" manifest:version="1.3"/>
 <manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>
 <manifest:file-entry manifest:full-path="styles.xml" manifest:media-type="text/xml"/>
 <manifest:file-entry manifest:full-path="meta.xml" manifest:media-type="text/xml"/>
 <manifest:file-entry manifest:full-path="settings.xml" manifest:media-type="text/xml"/>
 <manifest:file-entry manifest:full-path="Pictures/object-type-chart.png" manifest:media-type="image/png"/>
</manifest:manifest>"""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(output_path, "w") as archive:
        archive.writestr("mimetype", mimetype, compress_type=zipfile.ZIP_STORED)
        archive.writestr("content.xml", content, compress_type=zipfile.ZIP_DEFLATED)
        archive.writestr("styles.xml", styles, compress_type=zipfile.ZIP_DEFLATED)
        archive.writestr("meta.xml", meta, compress_type=zipfile.ZIP_DEFLATED)
        archive.writestr("settings.xml", settings, compress_type=zipfile.ZIP_DEFLATED)
        archive.writestr("META-INF/manifest.xml", manifest, compress_type=zipfile.ZIP_DEFLATED)
        archive.write(figure_path, "Pictures/object-type-chart.png", compress_type=zipfile.ZIP_DEFLATED)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--repo-root", type=Path, required=True)
    args = parser.parse_args()
    root = args.repo_root.resolve()
    output_dir = root / "examples" / "scholarly-work-foundations" / "output"
    audit_dir = root / "examples" / "scholarly-work-foundations" / "citation-style-audit"
    qa_dir = root / ".codex-tmp" / "issue-20-documents"
    output_dir.mkdir(parents=True, exist_ok=True)
    audit_dir.mkdir(parents=True, exist_ok=True)
    qa_dir.mkdir(parents=True, exist_ok=True)
    figure_path = qa_dir / "object-type-chart.png"
    build_figure(figure_path)
    build_structured_paper(output_dir / "structured-paper.docx", figure_path)
    build_odt(output_dir / "structured-paper.odt", figure_path)
    build_bibliography(audit_dir / "generated-bibliography.docx")
    print(output_dir / "structured-paper.docx")
    print(output_dir / "structured-paper.odt")
    print(audit_dir / "generated-bibliography.docx")


if __name__ == "__main__":
    main()
